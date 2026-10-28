#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把 BigMap **运行时 13 张表**导成策划配表 Excel，落到 `doc/塔防数据配置`。

与 `gen_bigmap_xlsx.py` 的分工（别搞混）：

| 脚本 | 内容 | 用途 |
| --- | --- | --- |
| `gen_bigmap_xlsx.py` | 按《植物大地图系统》**文档字段稿**写的 26 张（含 8 张存档表），字段名与运行时不一致 | 将来 P1/P2 扩表时的需求底稿 |
| **本脚本** | 只导**客户端真正在读**的 13 张，字段名与 `BigMapConfig.parseBigMapConfig` **逐字对齐** | 给策划改数值，走 `config:export` 直接覆盖 remoteRef |

格式严格照 `D:/City1001/doc/塔防数据配置` 现有 86 张表的六行表头：

    第 1 行  表名（<json 名>.json，合并整行）
    第 2 行  字段端（1=服务器+客户端 3=仅客户端 0=策划备注）
    第 3 行  字段中文名
    第 4 行  字段说明
    第 5 行  字段英文名        ← 必须等于 BigMapConfig 里 row['xxx'] 的那个 xxx
    第 6 行  数据类型（int / float / string）
    第 7 行起 数据行

数据行**直接取自 `gen_bigmap_tables.TABLES`**（同一个 `write()` 写的），
所以「游戏内 JSON」与「策划 Excel」永远同一套数字，不存在两处维护。

用法：
    python tools/config/gen_bigmap_esheets.py                # 默认写 doc/塔防数据配置
    python tools/config/gen_bigmap_esheets.py <输出目录> [--force]

已存在的同名 xlsx **默认跳过**（策划改过的表不能被脚本覆盖），要覆盖加 `--force`。
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))

import gen_bigmap_tables as tables  # noqa: E402

DEFAULT_OUT = Path(r'D:/City1001/doc/塔防数据配置')

# ── 样式（与现有策划表观感一致：深蓝表名行 + 灰蓝字段行 + 浅黄中文名行） ──────
FILL_TITLE = PatternFill('solid', fgColor='2F5597')
FILL_ROW2 = PatternFill('solid', fgColor='D9E2F3')
FILL_ROW3 = PatternFill('solid', fgColor='FFF2CC')
FILL_ROW4 = PatternFill('solid', fgColor='F2F2F2')
FONT_TITLE = Font(color='FFFFFF', bold=True, size=11)
FONT_BODY = Font(size=10)
CENTER = Alignment(horizontal='center', vertical='center')
WRAP = Alignment(horizontal='left', vertical='center', wrap_text=True)


def F(name: str, title: str, desc: str, side: int = 1, kind: str = 'int'):
    """一个字段：英文名 / 中文名 / 说明 / 端 / 类型。"""
    return {'name': name, 'title': title, 'desc': desc, 'side': side, 'kind': kind}


def S(name: str, title: str, desc: str, side: int = 1):
    """字符串字段（省得每次手写 kind）。"""
    return F(name, title, desc, side, 'string')


# ── 13 张表的字段元数据 ────────────────────────────────────────────────────
# 顺序必须与 `gen_bigmap_tables.write()` 的 key 列表**完全一致**，脚本会强校验。
META: dict[str, list[dict]] = {
    'bigmapWorld': [
        F('worldId', '世界ID', '大地图世界唯一 id，联机版由服务器下发；换 id 等于换一张图'),
        S('worldName', '世界名称', '显示用名称'),
        F('rows', '地图行数', '六边格行数（当前 40）'),
        F('cols', '地图列数', '六边格列数（当前 40）'),
        F('cityRing1', '个人层城池环', '个人层城池所在环半径，必须落在 (2, 地图半径) 内'),
        F('cityRing2', '军团层城池环', '军团层（本期不开放）',
          3),
        F('baseRadiusMin', '基地分布最小半径', '玩家/NPC 基地落位的最近环'),
        F('baseRadiusMax', '基地分布最大半径', '玩家/NPC 基地落位的最远环，不得小于最小值'),
        F('npcBaseCount', 'NPC 基地数量', '单机档模拟其他玩家的基地个数'),
        F('monsterRingMax', '野怪最大环', '野怪落位环带上限'),
    ],
    'bigmapSchedule': [
        F('scheduleId', '日程ID', '日程配置唯一 id'),
        F('utcOffsetHours', '时区偏移(小时)', '日程时间按哪个时区解释；东八区填 8'),
        S('openTime', '开放时间', '开放宣战/攻城，格式 HH:MM'),
        S('lockTime', '锁榜时间', '停止计分（但此时仍可攻城），格式 HH:MM，不得早于开放时间'),
        S('settleTime', '结算时间', '归属结算，格式 HH:MM，不得早于锁榜时间'),
        S('snapshotTime', '快照时间', '快照发奖，之后进入休战，格式 HH:MM，不得早于结算时间'),
    ],
    'bigmapBuilding': [
        F('id', '建筑ID', '建筑唯一 id（主城 71001、粮食田 71002…）'),
        S('baseBuildName', '建筑名称', '建造列表与详情面板显示名'),
        F('baseBuildType', '建筑类型', '1=主城 2=资源田 3=货币田 4=科技 5=建造队列 6=装饰'),
        F('occupyHex', '占用格数', '在建造区占用的六边格数'),
        F('maxOwnNum', '拥有上限', '每名玩家可同时拥有的数量'),
        F('unlockChapter', '解锁章节', '通关冒险第 N 章解锁（0=默认解锁）'),
        F('queueType', '队列类型', '1=建造队列 2=研发队列'),
        F('sort', '排序权重', '建造列表排序，越小越靠前'),
        S('baseBuildInfo', '建筑说明', '详情面板说明文案', 3),
    ],
    'bigmapBuildingLevel': [
        F('id', '建筑ID', '外键，关联 bigmapBuilding.id（建筑配置）'),
        F('buildLevel', '建筑等级', '必须从 1 连续排到最高级，中间不许缺'),
        F('costGrain', '消耗粮食', '升级消耗粮食'),
        F('costWood', '消耗木材', '升级消耗木材'),
        F('costIron', '消耗铁矿', '升级消耗铁矿'),
        F('instantCost', '立即完成钻石', '花钻石秒完成升级的价格'),
        F('preBuildId', '前置建筑ID', '升级前置建筑 id（0=无前置）'),
        F('preBuildLevel', '前置建筑等级', '前置建筑需要达到的等级'),
        F('buildTime', '建造耗时(秒)', '升级耗时；0 表示瞬发'),
        F('maxOtherLevel', '其他建筑等级上限', '主城专用：该等级下其他建筑的等级上限'),
        F('resCap', '资源存储上限', '主城专用：该等级下三种建造资源的存储上限'),
        F('attrBonus', '属性加成千分比', '主城伤害加成，千分比（50 = +5%）'),
    ],
    'bigmapBuildingProduce': [
        F('id', '建筑ID', '外键，关联 bigmapBuilding.id（建筑配置）'),
        F('buildLevel', '建筑等级', '产出档位'),
        S('resType', '产出资源', 'grain / wood / iron / coin，写错会被 parse 当场拒绝'),
        F('producePerHour', '每小时产量', '该等级的每小时产量', 1, 'float'),
        F('produceCap', '累积上限', '离线可累积的上限，不得小于每小时产量', 1, 'float'),
        S('techGroup', '科技组', '受哪类科技加成（economy/war/season），本期未接科技树', 3),
    ],
    'bigmapResource': [
        S('resType', '资源类型', 'grain / wood / iron / coin / stamina，五种必须齐'),
        S('resName', '资源名称', '资源栏显示名（与背包道具名保持一致）'),
        S('icon', '图标路径', '相对 assets/resources，不带扩展名', 3),
        F('startAmount', '初始数量', '新档初始值'),
        F('showInBar', '是否进资源栏', '1=进顶栏 0=不进（行动力走独立显示）', 3),
        F('itemId', '道具ID', '背包道具 id（18501..18505）：图标优先走 item.icon，0 = 只用概念图'),
    ],
    'bigmapQueue': [
        F('id', '配置ID', '队列配置唯一 id'),
        F('maxBuildQueue', '建造队列上限', '硬上限'),
        F('initBuildQueue', '初始建造队列', '新档默认拥有，不得大于上限'),
        F('queueUnlockCost2', '第2队列解锁价', '解锁第 2 条队列需要的货币', 1, 'float'),
        F('queueUnlockCost3', '第3队列解锁价', '解锁第 3 条队列需要的货币', 1, 'float'),
        F('cancelReturnRate', '取消返还比例', '0~1，1=全额返还', 1, 'float'),
        F('offlineMaxHours', '离线结算上限(小时)', '离线收益最多结算这么久'),
        F('upgradeUnlockChapter', '升级解锁章节', '第 N 章解锁建筑升级'),
        F('plotRows', '建造区行数', '必须是奇数（中心格放主城）'),
        F('plotCols', '建造区列数', '必须是奇数'),
    ],
    'bigmapCity': [
        F('cityId', '城池ID', '城池唯一 id'),
        S('cityName', '城池名称', '地图与详情面板显示名'),
        F('layer', '层级', '1=个人 2=军团 3=战盟 4=王座；本期只有 1 可开放'),
        F('available', '本期开放', '1=开放 0=画灰；开放层必须是个人层'),
        F('cityLevel', '城池等级', '外键，关联 bigmapCityLevel.cityLevel（城池等级）'),
        F('exclusive', '独占', '1=独占；本期客户端未读取，但服务端要用，端标记必须是 1（写成 0 会被 eSon 直接丢掉该列）'),
        F('unlockChapter', '解锁章节', '通关冒险第 N 章可打'),
        F('openDay', '开服天数', '开服第 N 天开放（0=当天）'),
        F('staminaCost', '行动力消耗', '填 0 则回退到等级表的值'),
        F('defenseHp', '城防', '填 0 则回退到等级表的值'),
        F('radiateCapacity', '辐射值', '⚠️口径待策划确认：当前代码当「势力范围半径(格)」用，文档另有「每城辐射 8~12 基地」的容量含义'),
        F('cityLevelCoeff', '等级系数', '填 0 则回退到等级表的值', 1, 'float'),
        F('prevCityId', '前置城池ID', '必须**占领**前一座才解锁（光打穿不算）；0=无前置'),
        F('rewardCoin', '占领奖励货币', '占领时一次性发放'),
        S('cityDesc', '城池说明', '详情面板文案', 3),
        F('slot', '槽位序号', '在所在环上的第几个位置，从 0 开始'),
        F('slotCount', '槽位总数', '所在环的槽位总数，slot 必须小于它'),
    ],
    'bigmapCityLevel': [
        F('cityLevel', '城池等级', '1~5'),
        F('warScoreRate', '战功倍率', '战功公式里的等级乘数', 1, 'float'),
        F('occupyScore', '占领积分', '占领时计入总积分'),
        F('powerThreshold', '上榜战力门槛', '战力不达标不进入榜单'),
        F('cityLevelCoeff', '等级系数', '战功公式里的等级系数', 1, 'float'),
        F('staminaCost', '行动力消耗', '默认消耗，城池表可覆盖'),
        F('minOccupyWarScore', '最低占领战功', '城防打穿后累计战功需达到它才能占领'),
        F('defenseHp', '城防', '默认城防，城池表可覆盖'),
    ],
    'bigmapRank': [
        F('rank', '名次', '必须从 1 连续排列'),
        F('coeffPct', '系数(百分比)', '100=不打折；第 11 名起按积分规则表的步长递减', 1, 'float'),
    ],
    'bigmapScoreRule': [
        F('scoreRuleId', '规则ID', '唯一 id'),
        F('marginalDamageCap', '伤害边际递减上限', '战功 = cap×(1-e^(-伤害/cap))'),
        F('rankExtendStepPct', '名次外延步长%', '第 11 名起每降一名扣多少个百分点', 1, 'float'),
        F('rankExtendMinPct', '名次系数下限%', '再差也不会低于它，不得超过 100', 1, 'float'),
        F('crowdBase', '拥挤系数基数', '0.5 = 人再多单次战功最低压到 50%', 1, 'float'),
        F('crowdK', '拥挤系数衰减率', '越大衰减越快', 1, 'float'),
        F('offRangeDamageCutPct', '越界伤害削减%', '基地不在势力范围时的伤害削减', 1, 'float'),
        F('npcScoreRatio', 'NPC战功比例', 'NPC 模拟榜的战功系数', 1, 'float'),
    ],
    'bigmapMonster': [
        F('monsterId', '野怪ID', '野怪唯一 id'),
        S('monsterName', '野怪名称', '地图与详情面板显示名'),
        F('monsterLevel', '野怪等级', '1~5'),
        F('occupyHex', '占用格数', '占用的六边格数'),
        F('preLevel', '前置等级', '需先通关该等级野怪才解锁（0=无门槛）'),
        F('staminaCost', '行动力消耗', '每次挑战消耗'),
        F('warScore', '战功', '击杀获得的战功'),
        F('defenseHp', '血量', '必须大于 0'),
        F('respawnSeconds', '刷新间隔(秒)', '击杀后多久重生'),
        F('rewardGrain', '掉落粮食', '击杀奖励'),
        F('rewardWood', '掉落木材', '击杀奖励'),
        F('rewardIron', '掉落铁矿', '击杀奖励'),
        F('rewardCoin', '掉落货币', '击杀奖励'),
    ],
    'bigmapStamina': [
        F('configId', '配置ID', '唯一 id'),
        F('staminaMax', '行动力上限', '必须大于 0'),
        F('recoverInterval', '恢复间隔(秒)', '多少秒恢复 1 点'),
        F('buyTimesDaily', '每日购买次数', '每日可买行动力的次数'),
        F('buyCost', '购买价格', '每次购买花费的货币', 1, 'float'),
        F('overflowKeep', '溢出保留', '上限后可溢出保留多少', 3),
    ],
}


def write_book(out_dir: Path, name: str, info: dict, force: bool) -> str:
    title = info['title']
    keys = info['keys']
    rows = info['rows']
    meta = META.get(name)
    if meta is None:
        raise KeyError(f'{name} 没有字段元数据（新增表要同时补 META）')
    if len(meta) != len(keys):
        raise ValueError(f'{name} 元数据 {len(meta)} 个字段，但表头有 {len(keys)} 个')
    for index, (field, key) in enumerate(zip(meta, keys)):
        # 元数据与 write() 的 key 顺序错位 ⇒ xlsx 里「中文名对、英文名串列」，策划必填错。
        if field['name'] != key:
            raise ValueError(f'{name} 第 {index + 1} 个字段错位：元数据 {field["name"]} vs 表头 {key}')

    out_path = out_dir / f'DM-{title}.xlsx'
    if out_path.exists() and not force:
        return f'跳过（已存在）{out_path.name}'

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]

    # 第 1 行：json 表名，合并整行
    ws.cell(1, 1).value = f'{name}.json'
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(keys)))
    ws.cell(1, 1).fill = FILL_TITLE
    ws.cell(1, 1).font = FONT_TITLE
    ws.cell(1, 1).alignment = CENTER

    for column, field in enumerate(meta, start=1):
        ws.cell(2, column).value = field['side']
        ws.cell(3, column).value = field['title']
        ws.cell(3, column).font = Font(size=10, bold=True)
        ws.cell(4, column).value = field['desc']
        ws.cell(5, column).value = field['name']
        ws.cell(5, column).font = Font(size=10, color='1F4E79')
        ws.cell(6, column).value = field['kind']
        ws.cell(2, column).fill = FILL_ROW2
        ws.cell(3, column).fill = FILL_ROW3
        ws.cell(4, column).fill = FILL_ROW4
        for row in (2, 4, 6):
            ws.cell(row, column).font = FONT_BODY
        for row in (2, 3, 5, 6):
            ws.cell(row, column).alignment = CENTER
        # 说明列最长，给 wrap + 按文本长度估宽（策划要能一眼读完）
        ws.cell(4, column).alignment = WRAP
        ws.column_dimensions[get_column_letter(column)].width = max(10, min(40, len(field['desc']) * 1.6))

    for row_index, values in enumerate(rows, start=7):
        for column, value in enumerate(values, start=1):
            ws.cell(row_index, column).value = value
            ws.cell(row_index, column).alignment = CENTER

    ws.freeze_panes = 'A7'
    ws.row_dimensions[4].height = 42
    thin = Side(style='thin', color='BFBFBF')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=2, max_row=6 + len(rows), max_col=max(1, len(keys))):
        for cell in row:
            cell.border = border

    wb.save(out_path)
    return f'新建 {out_path.name}（{len(rows)} 行 × {len(keys)} 列）'


def main() -> int:
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    force = '--force' in sys.argv
    out_dir = Path(args[0]) if args else DEFAULT_OUT
    out_dir.mkdir(parents=True, exist_ok=True)

    # 复用 gen_bigmap_tables 的 write()：跑一遍 main()，表数据就落在 TABLES 里
    tables.main()
    if not tables.TABLES:
        raise SystemExit('TABLES 为空：gen_bigmap_tables.main() 没有写出任何表')

    print(f'\n→ 输出目录 {out_dir}')
    created = 0
    for name, info in tables.TABLES.items():
        result = write_book(out_dir, name, info, force)
        print(f"  {info['title']}（{name}.json）: {result}")
        if result.startswith('新建'):
            created += 1
    print(f'\n共 {len(tables.TABLES)} 张表，新建 {created} 张，跳过 {len(tables.TABLES) - created} 张')
    return 0


if __name__ == '__main__':
    sys.exit(main())
