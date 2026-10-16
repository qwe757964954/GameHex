import openpyxl

path = r'D:\City1001\doc\塔防数据配置\S-商城.xlsx'
wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
print('sheets:', wb.sheetnames)
ws = wb['首充（废弃）']
for row in ws.iter_rows():
    for c in row:
        if c.value is not None:
            print(f'{c.coordinate}: {repr(c.value)}')
wb.close()
