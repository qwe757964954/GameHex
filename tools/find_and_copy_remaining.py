#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
从 FairyGUI 源目录里搜索剩余图标，复制到工程，再更新 Excel
"""

import os
import shutil
import openpyxl

UI_ICON_ROOT = r"D:\City1001\zarts\art\UIProject\trunk\UIIcon\assets"
PROJECT_ICONS = r"D:\City1001\Demo3.8\assets\resources\icons"
ITEM_TABLE = r"D:\City1001\doc\塔防数据配置\W-物品表.xlsx"


def classify_icon(icon_name):
    name = icon_name.lower()
    if "宝石" in icon_name or "gem" in name:
        return "gem"
    elif "装备" in icon_name or "equip" in name or any(k in name for k in ["wq", "ss", "kz", "mz", "jz", "xz", "yf", "xil", "xiz"]):
        return "equip"
    elif "头像" in icon_name or "head" in name or "hero_head" in name:
        return "hero"
    elif "技能" in icon_name or "skill" in name or "主公技" in icon_name:
        return "skill"
    elif "兵书" in icon_name or "兵种" in icon_name:
        return "item"
    elif any(k in icon_name for k in ["金币", "铜币", "元宝", "体力", "钻石", "灵石", "经验"]):
        return "resource"
    elif any(k in icon_name for k in ["宝箱", "钥匙", "卷轴", "图纸", "材料", "碎片", "兵书"]):
        return "item"
    else:
        return "common"


def main():
    print("=" * 60)
    print("从 FairyGUI 源目录搜索并复制剩余图标")
    print("=" * 60)

    # 1. 收集源目录所有 png
    print("\n[1] 扫描 FairyGUI 源目录所有 png...")
    all_source_pngs = {}  # 文件名(不带扩展) -> 完整路径
    for root, dirs, files in os.walk(UI_ICON_ROOT):
        for f in files:
            if f.endswith(".png"):
                name = f.replace(".png", "")
                full_path = os.path.join(root, f)
                all_source_pngs[name] = full_path

    print(f"  源目录共 {len(all_source_pngs)} 个 png")

    # 2. 打开 Excel，找待处理的
    print("\n[2] 读取 Excel 待处理物品...")
    wb = openpyxl.load_workbook(ITEM_TABLE)
    ws = wb["物品"]

    pending = []
    for row_idx in range(7, ws.max_row + 1):
        item_id_val = ws.cell(row=row_idx, column=1).value
        if not item_id_val:
            continue
        try:
            item_id = int(float(item_id_val))
        except (ValueError, TypeError):
            continue

        icon_val = ws.cell(row=row_idx, column=8).value
        if not icon_val:
            continue
        icon_str = str(icon_val).strip()
        if icon_str.startswith("icons/"):
            continue

        name = ws.cell(row=row_idx, column=2).value
        pending.append((row_idx, item_id, str(name), icon_str))

    print(f"  待处理: {len(pending)} 个")

    # 3. 逐个搜索、复制、更新
    copied_icons = {}  # icon_key -> 工程路径
    updated_rows = 0
    still_missing = []

    for row_idx, item_id, name, icon_key in pending:
        src_file = None

        # 精确匹配文件名
        if icon_key in all_source_pngs:
            src_file = all_source_pngs[icon_key]

        # 模糊匹配：去掉"图标_"前缀
        if not src_file:
            stripped = icon_key.replace("图标_", "")
            if stripped in all_source_pngs:
                src_file = all_source_pngs[stripped]

        # 更宽松的包含匹配
        if not src_file:
            for src_name, src_path in all_source_pngs.items():
                if icon_key in src_name or src_name in icon_key:
                    src_file = src_path
                    break

        if not src_file:
            still_missing.append((item_id, name, icon_key))
            continue

        # 分类并复制
        category = classify_icon(icon_key)
        cat_dir = os.path.join(PROJECT_ICONS, category)
        os.makedirs(cat_dir, exist_ok=True)

        filename = os.path.basename(src_file)
        dst_file = os.path.join(cat_dir, filename)

        if not os.path.exists(dst_file):
            shutil.copy2(src_file, dst_file)

        # 工程路径
        file_base = filename.replace(".png", "")
        proj_path = f"icons/{category}/{file_base}"
        copied_icons[icon_key] = proj_path

        # 更新 Excel
        ws.cell(row=row_idx, column=8, value=proj_path)
        updated_rows += 1

    wb.save(ITEM_TABLE)
    wb.close()

    print(f"\n[3] 结果:")
    print(f"  复制并更新: {updated_rows} 个")
    print(f"  仍未找到: {len(still_missing)} 个")

    if still_missing:
        print("\n=== 仍未找到的 ===")
        for item_id, name, icon_key in still_missing:
            print(f"  {item_id} {name}: {icon_key}")

    print("\n" + "=" * 60)
    print("完成！")
    print("=" * 60)


if __name__ == "__main__":
    main()
