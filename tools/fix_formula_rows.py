#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
修复 Excel 里的公式行：
1. A 列 ID 公式 → 实际数字
2. B 列名称公式 → 实际名称
3. H 列 icon → 工程路径（按品质 wq1/wq2/wq3）
"""

import openpyxl

ITEM_TABLE = r"D:\City1001\doc\塔防数据配置\W-物品表.xlsx"

# 起始：行169 = 14011（1阶精良）
START_ROW = 169
START_ID = 14011

QUALITY_NAMES = ["精良", "优秀", "史诗"]
QUALITY_ICONS = [
    "icons/equip/wq1",
    "icons/equip/wq2",
    "icons/equip/wq3",
]


def main():
    print("=" * 60)
    print("修复公式行：ID/名称/icon")
    print("=" * 60)

    wb = openpyxl.load_workbook(ITEM_TABLE)
    ws = wb["物品"]

    fixed = 0
    for row_idx in range(START_ROW, ws.max_row + 1):
        a_val = ws.cell(row=row_idx, column=1).value

        # 只处理 A 列是公式的行
        if not a_val or not str(a_val).startswith("="):
            continue

        # 计算实际 ID
        offset = row_idx - START_ROW
        tier_offset = offset // 3
        quality_idx = offset % 3
        actual_id = START_ID + tier_offset * 10 + quality_idx

        # 计算实际名称
        tier_num = tier_offset + 1
        quality_name = QUALITY_NAMES[quality_idx]
        actual_name = f"{tier_num}阶{quality_name}装备"

        # 写入 A 列：实际 ID
        ws.cell(row=row_idx, column=1, value=actual_id)

        # 写入 B 列：实际名称（不管原来是不是公式）
        ws.cell(row=row_idx, column=2, value=actual_name)

        # 写入 H 列：icon 工程路径
        ws.cell(row=row_idx, column=8, value=QUALITY_ICONS[quality_idx])

        fixed += 1
        if fixed <= 10 or fixed % 20 == 0:
            print(f"  行{row_idx}: {actual_id} {actual_name} -> {QUALITY_ICONS[quality_idx]}")

    wb.save(ITEM_TABLE)
    wb.close()

    print(f"\n修复完成: {fixed} 行")
    print("=" * 60)


if __name__ == "__main__":
    main()
