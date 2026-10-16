#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
把 Excel 物品表所有物品的 icon 字段同步成工程实际路径
流程：
1. 解析 T-图标.xlsx：icon key → FairyGUI 路径
2. 解析 FairyGUI 包：FairyGUI 路径 → 实际 png 文件
3. 按分类规则映射到工程 icons/ 目录下的路径
4. 写回 Excel 物品表的 icon 字段
"""

import os
import shutil
import xml.etree.ElementTree as ET
import openpyxl

# ============ 路径配置 ============
UI_ICON_ROOT = r"D:\City1001\zarts\art\UIProject\trunk\UIIcon\assets"
CONFIG_DIR = r"D:\City1001\doc\塔防数据配置"
ICON_TABLE = CONFIG_DIR + r"\T-图标.xlsx"
ITEM_TABLE = CONFIG_DIR + r"\W-物品表.xlsx"


def parse_fairy_gui_packages():
    """遍历 UIIcon/assets 下所有含 package.xml 的目录"""
    packages = {}
    for entry in os.listdir(UI_ICON_ROOT):
        pkg_dir = os.path.join(UI_ICON_ROOT, entry)
        pkg_xml = os.path.join(pkg_dir, "package.xml")
        if not os.path.isdir(pkg_dir) or not os.path.exists(pkg_xml):
            continue
        pkg_name = entry
        packages[pkg_name] = {}
        try:
            tree = ET.parse(pkg_xml)
            root = tree.getroot()
            for img in root.findall(".//image"):
                name_attr = img.get("name", "")
                path_attr = img.get("path", "")
                res_name = name_attr.replace(".png", "")
                rel_path = path_attr + name_attr if path_attr else name_attr
                rel_path = rel_path.replace("/", os.sep).lstrip(os.sep)
                packages[pkg_name][res_name] = rel_path
        except Exception as e:
            print(f"  [错误] 解析 {pkg_xml} 失败: {e}")
    return packages


def parse_icon_table():
    """读取 T-图标.xlsx，返回 {key: fairy_gui_path}"""
    icon_map = {}
    wb = openpyxl.load_workbook(ICON_TABLE, read_only=True)
    ws = wb["Sheet1"]
    for row in ws.iter_rows(min_row=7, values_only=True):
        key, path = row[0], row[1]
        if key and path and str(path).startswith("ui://"):
            icon_map[str(key)] = str(path)
    wb.close()
    return icon_map


def classify_icon(icon_name, fairy_path=""):
    """根据图标名判断分类"""
    name = icon_name.lower()
    if "宝石" in icon_name or "gem" in name:
        return "gem"
    elif "装备" in icon_name or "equip" in name or any(k in name for k in ["wq", "ss", "kz", "mz", "jz", "xz", "yf", "xil", "xiz"]):
        return "equip"
    elif "头像" in icon_name or "head" in name or "hero_head" in name:
        return "hero"
    elif "技能" in icon_name or "skill" in name:
        return "skill"
    elif "buff" in name:
        return "buff"
    elif any(k in icon_name for k in ["金币", "铜币", "元宝", "体力", "钻石", "灵石", "经验"]):
        return "resource"
    elif any(k in icon_name for k in ["宝箱", "钥匙", "卷轴", "图纸", "材料", "碎片"]):
        return "item"
    else:
        return "common"


def main():
    print("=" * 60)
    print("同步 Excel 物品表 icon 字段为工程路径")
    print("=" * 60)

    # 1. 解析 FairyGUI 包
    print("\n[1] 解析 FairyGUI 图标包...")
    packages = parse_fairy_gui_packages()
    print(f"  {len(packages)} 个包")

    # 2. 解析 T-图标表
    print("\n[2] 解析 T-图标.xlsx...")
    icon_table = parse_icon_table()
    print(f"  {len(icon_table)} 条映射")

    # 3. 建立 icon key -> 工程路径 的映射
    print("\n[3] 建立 icon key -> 工程路径...")

    def resolve_icon_path(icon_key):
        fairy_path = icon_table.get(icon_key)
        if not fairy_path:
            return None

        parts = fairy_path.replace("ui://", "").split("/")
        if len(parts) != 2:
            return None

        pkg_name, res_name = parts[0], parts[1]
        if pkg_name not in packages:
            return None

        pkg_resources = packages[pkg_name]

        # 精确匹配
        if res_name in pkg_resources:
            rel_path = pkg_resources[res_name]
            filename = os.path.basename(rel_path).replace(".png", "")
            category = classify_icon(icon_key)
            return f"icons/{category}/{filename}"

        # 模糊匹配
        for k, v in pkg_resources.items():
            if res_name in k or k in res_name:
                rel_path = v
                filename = os.path.basename(rel_path).replace(".png", "")
                category = classify_icon(icon_key)
                return f"icons/{category}/{filename}"

        return None

    icon_key_to_path = {}
    for key in icon_table:
        path = resolve_icon_path(key)
        if path:
            icon_key_to_path[key] = path

    print(f"  成功映射: {len(icon_key_to_path)} 个")

    # 4. 更新 Excel 物品表
    print("\n[4] 更新 W-物品表.xlsx...")
    wb = openpyxl.load_workbook(ITEM_TABLE)
    ws = wb["物品"]

    updated = 0
    skipped = 0
    not_found = 0

    for row_idx in range(7, ws.max_row + 1):
        item_id_val = ws.cell(row=row_idx, column=1).value
        if not item_id_val:
            continue

        try:
            item_id = int(float(item_id_val))
        except (ValueError, TypeError):
            continue

        old_icon = ws.cell(row=row_idx, column=8).value
        if not old_icon:
            continue

        old_icon_str = str(old_icon).strip()

        # 已经是路径格式的（以 icons/ 开头），跳过
        if old_icon_str.startswith("icons/"):
            skipped += 1
            continue

        # 查映射
        if old_icon_str in icon_key_to_path:
            new_path = icon_key_to_path[old_icon_str]
            ws.cell(row=row_idx, column=8, value=new_path)
            updated += 1
            if updated <= 5:
                name = ws.cell(row=row_idx, column=2).value
                print(f"  {item_id} {name}: {old_icon_str} -> {new_path}")
        else:
            not_found += 1

    wb.save(ITEM_TABLE)
    wb.close()

    print(f"\n  更新: {updated} 个")
    print(f"  已是路径(跳过): {skipped} 个")
    print(f"  未找到映射: {not_found} 个")

    print("\n" + "=" * 60)
    print("完成！")
    print("=" * 60)


if __name__ == "__main__":
    main()
