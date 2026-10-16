#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
补全剩余未匹配的图标：
1. 找出 Excel 里还没改成路径的物品
2. 直接在工程 icons 目录里做宽松匹配
3. 能找到的都改了
"""

import os
import openpyxl

ITEM_TABLE = r"D:\City1001\doc\塔防数据配置\W-物品表.xlsx"
PROJECT_ICONS = r"D:\City1001\Demo3.8\assets\resources\icons"


def main():
    print("=" * 60)
    print("补全剩余图标映射")
    print("=" * 60)

    # 1. 收集工程里所有图标文件
    all_project_icons = {}  # 文件名(不带扩展) -> 相对路径
    for root, dirs, files in os.walk(PROJECT_ICONS):
        for f in files:
            if f.endswith(".png"):
                name = f.replace(".png", "")
                rel_path = os.path.relpath(os.path.join(root, f), PROJECT_ICONS)
                rel_path = rel_path.replace(os.sep, "/").replace(".png", "")
                all_project_icons[name] = f"icons/{rel_path}"

    print(f"\n工程图标总数: {len(all_project_icons)}")

    # 2. 打开 Excel，找还没改的
    wb = openpyxl.load_workbook(ITEM_TABLE)
    ws = wb["物品"]

    pending = []
    for row_idx in range(7, ws.max_row + 1):
        item_id_val = ws.cell(row=row_idx, column=1).value
        if not item_id_val:
            continue
        try:
            item_id = int(float(item_id_val))
        except (ValueError, TypeError):
            continue

        icon_val = ws.cell(row=row_idx, column=8).value
        if not icon_val:
            continue

        icon_str = str(icon_val).strip()
        if icon_str.startswith("icons/"):
            continue  # 已经改了

        name = ws.cell(row=row_idx, column=2).value
        pending.append((row_idx, item_id, str(name), icon_str))

    print(f"待处理: {len(pending)} 个")
    print()

    # 3. 逐个匹配
    updated = 0
    still_missing = []

    for row_idx, item_id, name, icon_key in pending:
        matched_path = None

        # 策略1：精确匹配文件名
        if icon_key in all_project_icons:
            matched_path = all_project_icons[icon_key]

        # 策略2：去掉"图标_"前缀再匹配
        if not matched_path:
            stripped = icon_key.replace("图标_", "")
            if stripped in all_project_icons:
                matched_path = all_project_icons[stripped]

        # 策略3：模糊包含（工程文件名包含 icon key，或反过来）
        if not matched_path:
            for proj_name, proj_path in all_project_icons.items():
                if icon_key in proj_name or proj_name in icon_key:
                    matched_path = proj_path
                    break

        # 策略4：去掉"随机"等词再试
        if not matched_path:
            cleaned = icon_key.replace("随机", "").replace("图标_", "")
            for proj_name, proj_path in all_project_icons.items():
                if cleaned and (cleaned in proj_name or proj_name in cleaned):
                    matched_path = proj_path
                    break

        if matched_path:
            ws.cell(row=row_idx, column=8, value=matched_path)
            updated += 1
            print(f"  ✓ {item_id} {name}: {icon_key} -> {matched_path}")
        else:
            still_missing.append((item_id, name, icon_key))

    wb.save(ITEM_TABLE)
    wb.close()

    print(f"\n匹配成功: {updated} 个")
    print(f"仍未找到: {len(still_missing)} 个")

    if still_missing:
        print("\n=== 仍未找到的 ===")
        for item_id, name, icon_key in still_missing:
            print(f"  {item_id} {name}: {icon_key}")

    print("\n" + "=" * 60)
    print("完成！")
    print("=" * 60)


if __name__ == "__main__":
    main()
