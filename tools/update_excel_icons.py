#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
直接修改 W-物品表.xlsx：
装备物品的 icon 字段直接填工程资源路径（icons/equip/wqX）
不再经过 T-图标表映射
"""

import openpyxl

ITEM_TABLE = r"D:\City1001\doc\塔防数据配置\W-物品表.xlsx"

# 装备品质对应工程路径
QUALITY_ICON = {
    1: "icons/equip/wq1",  # 精良
    2: "icons/equip/wq2",  # 优秀
    3: "icons/equip/wq3",  # 史诗
    4: "icons/equip/wq4",  # 传说
    5: "icons/equip/wq4",
}


def main():
    print("=" * 60)
    print("修改 W-物品表.xlsx：装备 icon 直接填工程路径")
    print("=" * 60)
    
    wb = openpyxl.load_workbook(ITEM_TABLE)
    ws = wb["物品"]
    
    updated = 0
    for row_idx in range(7, ws.max_row + 1):
        item_id_val = ws.cell(row=row_idx, column=1).value
        if not item_id_val:
            continue
        
        try:
            item_id = int(float(item_id_val))
        except (ValueError, TypeError):
            continue
        
        id_str = str(item_id)
        # 只处理 14xxx 装备
        if not id_str.startswith("14") or len(id_str) != 5:
            continue
        
        # ID 结构: 14 XY Z
        # X Y = 阶数（01=1阶, 02=2阶...）
        # Z = 品质（最后一位）：1=精良, 2=优秀, 3=史诗
        quality_digit = int(id_str[4])  # 第五位是品质
        if quality_digit in QUALITY_ICON:
            new_path = QUALITY_ICON[quality_digit]
            old_icon = ws.cell(row=row_idx, column=8).value
            ws.cell(row=row_idx, column=8, value=new_path)
            updated += 1
            if updated <= 5:
                name = ws.cell(row=row_idx, column=2).value
                print(f"  {item_id} {name}: {old_icon} -> {new_path}")
    
    wb.save(ITEM_TABLE)
    wb.close()
    
    print(f"\n完成！更新了 {updated} 个装备物品的 icon 字段")
    print("=" * 60)


if __name__ == "__main__":
    main()
